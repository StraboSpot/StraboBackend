"""Answer sheets for the testers (step 6 answer-key points 1-3).

One .xlsx per tester per outing (local date of the recordings), made from the
server export AFTER the outing. Pre-filled ONLY with each Voice Spot's name
and recording start time (never anything proposed or saved: scoring stays
blind), one header row per Spot with its recording id in a hidden column.
The tester adds one row per measurement under each header (insert rows as
needed) and fills the Stopwatch tab.
"""

import datetime as dt
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

VOICE = 'Voice Spots'
STOPWATCH = 'Stopwatch'
LISTS = 'Lists'
CONVENTIONS = ['Right-hand rule', 'Dip direction always written']
KINDS = ['Plane', 'Line', 'Sample']
COLUMNS = ['Spot', 'Recorded at', '#', 'Kind', 'Feature type', 'Strike', 'Dip', 'Dip direction',
           'Trend', 'Plunge', 'Quality', 'Lies on #', 'Movement / facing', 'Sample label', 'Photos', 'As written',
           'Recording id']   # last column, hidden (rename-proof matching)
COL = {name: i + 1 for i, name in enumerate(COLUMNS)}
FIRST_ROW = 7          # header row of the table; Spots start below it
CONVENTION_CELL = 'C4'
BLANK_ROWS = 3         # measurement rows offered under each Spot
STOPWATCH_COLUMNS = ['Date', 'StraboField Spot name', 'Time (mm:ss)', 'Note']


def local_start(st):
    t = dt.datetime.strptime(st['started_at'][:19], '%Y-%m-%dT%H:%M:%S').replace(tzinfo=dt.timezone.utc)
    return t + dt.timedelta(minutes=st.get('tz_offset_minutes') or 0)


def feature_labels(vocab):
    """Planar + linear feature type labels (the app's own form choices)."""
    forms = (vocab or {}).get('forms') or {}
    labels = []
    for form in ('measurement.planar_orientation', 'measurement.linear_orientation'):
        for _name, label in ((forms.get(form) or {}).get('feature_type') or {}).items():
            if label not in labels:
                labels.append(label)
    return labels


def outings(export):
    """{(userpkey, 'YYYY-MM-DD'): [stations]} for confirmed recordings, in recording order."""
    out = {}
    for st in export['stations']:
        if st.get('outcome') != 'confirmed':
            continue
        out.setdefault((st['userpkey'], local_start(st).date().isoformat()), []).append(st)
    for v in out.values():
        v.sort(key=lambda s: s['started_at'])
    return out


def build(export, upk, day, stations):
    tester = next((t for t in export.get('testers') or [] if t['userpkey'] == upk), {})
    wb = Workbook()
    ws = wb.active
    ws.title = VOICE
    bold = Font(bold=True)
    ws['A1'] = f'Strabo Voice answer sheet: {tester.get("name") or tester.get("email")}, {day}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = ('From your NOTEBOOK only. One row per measurement under each Spot (insert rows as needed). '
                'Leave a cell blank if the notebook does not have it.')
    ws['A4'] = 'Strike convention in your notebook:'
    ws['A4'].font = bold
    ws[CONVENTION_CELL] = CONVENTIONS[0]
    ws['A5'] = 'Odd notation (quadrants like N45E 32SE)? Put it in "As written" exactly as in the notebook.'
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=FIRST_ROW, column=c, value=name)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical='top')
    head_fill = PatternFill('solid', fgColor='DDE7F3')
    r = FIRST_ROW + 1
    for st in stations:
        ws.cell(row=r, column=COL['Recording id'], value=st['station_uuid'])
        ws.cell(row=r, column=COL['Spot'], value=st.get('spot_name') or '')
        ws.cell(row=r, column=COL['Recorded at'], value=local_start(st).strftime('%H:%M'))
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=c).fill = head_fill
        r += 1 + BLANK_ROWS
    last = max(r + 50, 200)

    lists = wb.create_sheet(LISTS)
    lists['A1'], lists['B1'] = 'tester', upk
    lists['A2'], lists['B2'] = 'outing', day
    lists['A3'], lists['B3'] = 'format', 1
    for i, v in enumerate(KINDS, start=1):
        lists.cell(row=i, column=4, value=v)
    for i, v in enumerate(feature_labels(export.get('vocab')), start=1):
        lists.cell(row=i, column=5, value=v)
    for i, v in enumerate(CONVENTIONS, start=1):
        lists.cell(row=i, column=6, value=v)
    n_ft = max(1, len(feature_labels(export.get('vocab'))))
    lists.sheet_state = 'hidden'

    def letter(name):
        return ws.cell(row=1, column=COL[name]).column_letter

    def add(dv, name):
        ws.add_data_validation(dv)
        dv.add(f'{letter(name)}{FIRST_ROW + 1}:{letter(name)}{last}')

    conv = DataValidation(type='list', formula1=f'={LISTS}!$F$1:$F${len(CONVENTIONS)}', allow_blank=False)
    ws.add_data_validation(conv)
    conv.add(CONVENTION_CELL)
    add(DataValidation(type='list', formula1=f'={LISTS}!$D$1:$D${len(KINDS)}', allow_blank=True), 'Kind')
    add(DataValidation(type='list', formula1=f'={LISTS}!$E$1:$E${n_ft}', allow_blank=True), 'Feature type')
    for name in ('Strike', 'Dip direction', 'Trend'):
        add(DataValidation(type='decimal', operator='between', formula1='0', formula2='360', allow_blank=True,
                           error='0 to 360 degrees', showErrorMessage=True), name)
    for name in ('Dip', 'Plunge'):
        add(DataValidation(type='decimal', operator='between', formula1='0', formula2='90', allow_blank=True,
                           error='0 to 90 degrees', showErrorMessage=True), name)
    add(DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True,
                       error='Quality is 1 to 5', showErrorMessage=True), 'Quality')
    for name in ('#', 'Lies on #', 'Photos'):
        add(DataValidation(type='whole', operator='between', formula1='0', formula2='999', allow_blank=True,
                           error='A whole number', showErrorMessage=True), name)
    ws.column_dimensions[letter('Recording id')].hidden = True
    widths = {'Spot': 12, 'Recorded at': 10, '#': 5, 'Kind': 9, 'Feature type': 16, 'Movement / facing': 18,
              'Sample label': 14, 'As written': 26}
    for name, w in widths.items():
        ws.column_dimensions[letter(name)].width = w
    ws.freeze_panes = ws.cell(row=FIRST_ROW + 1, column=COL['#'])

    sw = wb.create_sheet(STOPWATCH)
    sw['A1'] = ('StraboField outcrops: start the stopwatch when you pick up the compass, stop it when the Spot is '
                'saved. Spots in the "Stopwatch Spots" dataset.')
    for c, name in enumerate(STOPWATCH_COLUMNS, start=1):
        sw.cell(row=3, column=c, value=name).font = bold
    for r in range(4, 200):
        sw.cell(row=r, column=3).number_format = '@'   # text: "1:35" stays minutes:seconds
    sw.column_dimensions['B'].width = 24
    sw.column_dimensions['D'].width = 30
    return wb


def safe(s):
    return re.sub(r'[^A-Za-z0-9]+', '_', str(s or '')).strip('_') or 'tester'


def write_all(export, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    testers = {t['userpkey']: t for t in export.get('testers') or []}
    for (upk, day), sts in sorted(outings(export).items()):
        t = testers.get(upk, {})
        path = os.path.join(out_dir, f'AnswerSheet_{safe(t.get("name") or t.get("email"))}_{day}.xlsx')
        build(export, upk, day, sts).save(path)
        paths.append(path)
    return paths
