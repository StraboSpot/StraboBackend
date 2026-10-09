"""Returned answer sheets + corrections log -> the trial answer key (answer-key point 4).

NEVER guesses: anything it cannot read cleanly goes on the problem list, and
scoring refuses to run until that list is empty. Quadrant notation in
"As written" is converted by code and listed for Jason to check
(`score --conversions-checked` once he has).

Key format (what scorer.py reads):
  {"format": 1,
   "recordings": {station_uuid: {"tester": upk, "spot_name": str, "photos": n|null,
       "measurements": [{"n", "kind": plane|line|sample, "ft": name|null, "strike", "dip",
           "dd", "trend", "plunge", "quality", "lies_on", "words", "sample", "corrected": [..]}]}},
   "baseline": [{"tester", "date", "spot_name", "seconds", "n_measurements", "note"}],
   "corrections": [the applied log lines], "conversions": [{.., "as_written", "dd"}]}

corrections.csv columns: date, tester (email), spot, n, field, old, new, reason, decided_by
  field = kind | ft | strike | dip | dd | trend | plunge | quality | lies_on | words | sample
"""

import csv
import datetime as dt
import glob
import hashlib
import os
import re

from openpyxl import load_workbook

from . import sheets as S
from .saved import num

FIELDS = {'Kind': 'kind', 'Feature type': 'ft', 'Strike': 'strike', 'Dip': 'dip', 'Dip direction': 'dd',
          'Trend': 'trend', 'Plunge': 'plunge', 'Quality': 'quality', 'Lies on #': 'lies_on',
          'Movement / facing': 'words', 'Sample label': 'sample', 'As written': 'as_written'}
CORRECTION_FIELDS = ('kind', 'ft', 'strike', 'dip', 'dd', 'trend', 'plunge', 'quality', 'lies_on', 'words', 'sample')


def key_sha256(paths):
    """One hash over every sheet + the corrections log (point 4)."""
    h = hashlib.sha256()
    for p in sorted(paths):
        h.update(os.path.basename(p).encode() + b'\0')
        with open(p, 'rb') as f:
            h.update(f.read())
        h.update(b'\0')
    return h.hexdigest()


def quadrant_strike(q1, deg, q2):
    deg = int(deg)
    return {('N', 'E'): deg, ('N', 'W'): (360 - deg) % 360, ('S', 'E'): 180 - deg, ('S', 'W'): 180 + deg}[(q1, q2)]


DIRS = {'N': (0, 1), 'S': (0, -1), 'E': (1, 0), 'W': (-1, 0)}


def dd_from(strike, dip_quadrant):
    """The dip direction (strike +- 90) on the side the dip quadrant names."""
    vx = sum(DIRS[c][0] for c in dip_quadrant)
    vy = sum(DIRS[c][1] for c in dip_quadrant)
    import math
    for cand in ((strike + 90) % 360, (strike - 90) % 360):
        r = math.radians(cand)
        if math.sin(r) * vx + math.cos(r) * vy > 0:
            return cand
    return None


def parse_written(s):
    """'N45E 32SE' or '045/32 SE' -> (strike, dip, dd) or None."""
    t = str(s or '').upper().replace(',', ' ')
    m = re.search(r'\b([NS])\s*(\d{1,2})\s*([EW])\b[\s/]*(\d{1,2})\s*([NSEW]{1,2})\b', t)
    if m:
        strike = quadrant_strike(m.group(1), m.group(2), m.group(3))
        dip = int(m.group(4))
    else:
        m = re.search(r'\b(\d{1,3})\s*[/ ]\s*(\d{1,2})\s*([NSEW]{1,2})\b', t)
        if not m or int(m.group(1)) > 360:
            return None
        strike, dip = int(m.group(1)) % 360, int(m.group(2))
    dd = dd_from(strike, m.groups()[-1])
    return None if dd is None or dip > 90 else (strike, dip, dd)


def _cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def parse_time(v):
    """Stopwatch time -> seconds: '1:35', '95', 95, an Excel time or duration."""
    if v is None or v == '':
        return None
    # A spreadsheet that turned a typed "1:35" into a clock time read it as
    # 1 h 35 min; the stopwatch meant 1 min 35 s.
    if isinstance(v, dt.timedelta):
        return v.total_seconds() / 60
    if isinstance(v, (dt.time, dt.datetime)):
        return v.hour * 60 + v.minute if v.second == 0 else v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, (int, float)):
        return round(float(v) * 1440, 1) if 0 < v < 1 else float(v)
    m = re.fullmatch(r'\s*(\d+):(\d{1,2})(?:\.(\d+))?\s*', str(v))
    if m:
        return int(m.group(1)) * 60 + int(m.group(2))
    try:
        return float(str(v))
    except ValueError:
        return None


def read_sheet(path, export, problems, conversions):
    wb = load_workbook(path, data_only=True)
    where = os.path.basename(path)
    lists = wb[S.LISTS]
    upk = int(lists['B1'].value)
    day = str(lists['B2'].value)
    vocab = (export.get('vocab') or {}).get('forms') or {}
    by_label = {}
    for form, kind in (('measurement.planar_orientation', 'plane'), ('measurement.linear_orientation', 'line')):
        for name, label in ((vocab.get(form) or {}).get('feature_type') or {}).items():
            by_label[(kind, str(label).lower())] = name
            by_label[(kind, str(name).lower())] = name
    ws = wb[S.VOICE]
    convention = _cell(ws[S.CONVENTION_CELL].value)
    if convention not in S.CONVENTIONS:
        problems.append(f'{where}: strike convention cell {S.CONVENTION_CELL} is "{convention}"')
    header = [str(c.value or '').strip() for c in ws[S.FIRST_ROW]]
    col = {name: header.index(name) for name in S.COLUMNS if name in header}
    for name in S.COLUMNS:
        if name not in col:
            problems.append(f'{where}: column "{name}" is missing')
            return {}
    recordings = {}
    current = None
    for row in ws.iter_rows(min_row=S.FIRST_ROW + 1, values_only=True):
        get = {name: _cell(row[i]) if i < len(row) else None for name, i in col.items()}
        if get['Recording id']:
            current = recordings.setdefault(str(get['Recording id']), {
                'tester': upk, 'spot_name': str(get['Spot'] or ''), 'photos': num(get['Photos']), 'measurements': []})
            continue
        vals = {FIELDS[k]: get[k] for k in FIELDS}
        if not any(v is not None for v in vals.values()):
            continue
        if current is None:
            problems.append(f'{where}: a measurement row above the first Spot')
            continue
        spot = current['spot_name']
        n = num(get['#']) or len(current['measurements']) + 1
        m = {'n': n, 'corrected': []}
        kind = str(vals['kind'] or '').lower()
        if kind not in ('plane', 'line', 'sample'):
            problems.append(f'{where} Spot {spot} #{n}: Kind "{vals["kind"]}" (Plane, Line or Sample)')
            continue
        m['kind'] = kind
        for f in ('strike', 'dip', 'dd', 'trend', 'plunge', 'quality', 'lies_on'):
            v = vals[f]
            m[f] = num(v)
            if v is not None and m[f] is None:
                problems.append(f'{where} Spot {spot} #{n}: {f} "{v}" is not a number')
        m['words'] = vals['words']
        m['sample'] = None if vals['sample'] is None else str(vals['sample'])
        m['ft'] = None
        if vals['ft'] is not None:
            m['ft'] = by_label.get((kind, str(vals['ft']).lower()))
            if m['ft'] is None:
                problems.append(f'{where} Spot {spot} #{n}: unknown feature type "{vals["ft"]}" for a {kind}')
        if kind == 'plane':
            if m['dd'] is None and vals['as_written']:
                p = parse_written(vals['as_written'])
                if p is None:
                    problems.append(f'{where} Spot {spot} #{n}: cannot read "As written" "{vals["as_written"]}"')
                else:
                    m['strike'] = m['strike'] if m['strike'] is not None else p[0]
                    m['dip'] = m['dip'] if m['dip'] is not None else p[1]
                    m['dd'] = p[2]
                    conversions.append({'sheet': where, 'spot': spot, 'n': n, 'as_written': vals['as_written'],
                                        'strike': p[0], 'dip': p[1], 'dd': p[2]})
            if m['dd'] is None and m['strike'] is not None and convention == S.CONVENTIONS[0]:
                m['dd'] = (m['strike'] + 90) % 360
            if m['dd'] is None and m['dip'] is not None:
                problems.append(f'{where} Spot {spot} #{n}: a plane without a dip direction '
                                f'(fill Dip direction, or Strike with the right-hand rule convention)')
        if kind == 'sample' and not m['sample']:
            problems.append(f'{where} Spot {spot} #{n}: a Sample row without a label')
        current['measurements'].append(m)
    for uuid, rec in recordings.items():
        ns = {m['n'] for m in rec['measurements']}
        if len(ns) != len(rec['measurements']):
            problems.append(f'{where} Spot {rec["spot_name"]}: two rows with the same #')
        for m in rec['measurements']:
            if m.get('lies_on') is not None:
                tgt = next((x for x in rec['measurements'] if x['n'] == m['lies_on']), None)
                if tgt is None or tgt['kind'] != 'plane':
                    problems.append(f'{where} Spot {rec["spot_name"]} #{m["n"]}: Lies on #{m["lies_on"]} is not a plane')
            if m['kind'] == 'line' and m.get('dd') is not None:
                problems.append(f'{where} Spot {rec["spot_name"]} #{m["n"]}: a Line with a dip direction')

    baseline = []
    sw = wb[S.STOPWATCH]
    stop_spots = {}
    for d in export.get('stopwatch') or []:
        if d['userpkey'] == upk:
            for s in d['spots']:
                stop_spots.setdefault(str(s.get('name') or '').strip().lower(), []).append(s)
    for row in sw.iter_rows(min_row=4, values_only=True):
        date, name, t, note = (list(row) + [None] * 4)[:4]
        if all(_cell(x) is None for x in (date, name, t, note)):
            continue
        secs = parse_time(t)
        name = str(_cell(name) or '')
        found = stop_spots.get(name.strip().lower()) or []
        if secs is None:
            problems.append(f'{where} Stopwatch "{name}": time "{t}" is not mm:ss')
        if len(found) != 1:
            problems.append(f'{where} Stopwatch "{name}": {"not found" if not found else "several"} '
                            f'in the tester\'s Stopwatch Spots datasets on the server')
            continue
        baseline.append({'tester': upk, 'date': str(_cell(date) or day), 'spot_name': name, 'seconds': secs,
                         'n_measurements': found[0]['n_measurements'], 'note': _cell(note)})
    return {'recordings': recordings, 'baseline': baseline, 'tester': upk}


def apply_corrections(path, key, export, problems):
    if not os.path.exists(path):
        return []
    emails = {(t.get('email') or '').lower(): t['userpkey'] for t in export.get('testers') or []}
    applied = []
    with open(path, newline='') as f:
        for i, row in enumerate(csv.DictReader(f), start=2):
            where = f'corrections.csv line {i}'
            upk = emails.get((row.get('tester') or '').strip().lower())
            field = (row.get('field') or '').strip()
            if upk is None or field not in CORRECTION_FIELDS:
                problems.append(f'{where}: unknown tester or field')
                continue
            recs = [r for r in key['recordings'].values()
                    if r['tester'] == upk and r['spot_name'] == (row.get('spot') or '').strip()]
            m = next((x for r in recs for x in r['measurements'] if x['n'] == num(row.get('n'))), None)
            if m is None:
                problems.append(f'{where}: no measurement {row.get("n")} on Spot {row.get("spot")}')
                continue
            old, new = row.get('old', ''), row.get('new', '')
            cur = m.get(field)
            if str('' if cur is None else cur) != str(num(old) if num(old) is not None else old).strip():
                problems.append(f'{where}: {field} is "{cur}" on the sheet, not "{old}"')
                continue
            m[field] = num(new) if field not in ('kind', 'ft', 'words', 'sample') else (new or None)
            m['corrected'].append(field)
            applied.append(dict(row))
    return applied


def convert(sheet_dir, export):
    """-> (key, problems, key_sha256)"""
    problems, conversions = [], []
    paths = sorted(glob.glob(os.path.join(sheet_dir, '*.xlsx')))
    key = {'format': 1, 'recordings': {}, 'baseline': [], 'corrections': [], 'conversions': conversions}
    for p in paths:
        part = read_sheet(p, export, problems, conversions)
        for uuid, rec in (part.get('recordings') or {}).items():
            if uuid in key['recordings']:
                problems.append(f'{os.path.basename(p)}: recording {uuid} is on two sheets')
            key['recordings'][uuid] = rec
        key['baseline'] += part.get('baseline') or []
    corr = os.path.join(sheet_dir, 'corrections.csv')
    key['corrections'] = apply_corrections(corr, key, export, problems)
    known = {st['station_uuid'] for st in export['stations']}
    for uuid in key['recordings']:
        if uuid not in known:
            problems.append(f'recording {uuid} on a sheet is not in the export')
    hashed = paths + ([corr] if os.path.exists(corr) else [])
    return key, problems, key_sha256(hashed)
